"""Copy owner-approved research files into the Excel register (one sheet per FORMAT.md section).

python research_to_excel.py <register.xlsx> <key> [<key> ...]

The research JSON is the source; the workbook is generated from it. Re-running replaces that exercise's rows
(so a corrected exercise never leaves stale rows behind). Only status "owner-approved" is accepted.
"""
import sys, os, json
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment

HERE = os.path.dirname(os.path.abspath(__file__))
RESEARCH = os.path.join(os.path.dirname(HERE), "research")

SHEETS = [  # sheet, section, columns (FORMAT.md order)
    ("Exercises", "exercise", ["key", "name", "equipment", "library_primary", "library_secondary", "status", "version",
                               "researched_on", "one_line_summary", "evidence_grade"]),
    ("Setup", "setup", ["key", "item", "target", "tolerance", "why", "source_ids", "grade"]),
    ("Movement", "movement", ["key", "joint", "measure", "start", "end", "path", "tolerance", "source_ids", "grade"]),
    ("Muscles", "muscles", ["key", "muscle", "anatomy", "mask_key", "tier", "glow", "emg", "role", "surface",
                            "source_ids", "grade"]),
    ("Tempo", "tempo", ["key", "up_s", "hold_top_s", "down_s", "pause_bottom_s", "loop_s", "source_ids"]),
    ("Mistakes", "mistakes", ["key", "mistake", "what it looks like", "why it matters", "source_ids"]),
    ("Sources", "sources", ["key", "id", "citation", "url", "type", "what_it_gave"]),
    ("Clip check", "clip_check", ["key", "check", "target", "measured", "pass", "tool"]),
]
HEAD = Font(bold=True, color="FFFFFF"); FILL = PatternFill("solid", fgColor="B4571B")


def cell(v):
    if isinstance(v, list):
        return ", ".join(str(x) for x in v)
    if isinstance(v, dict):
        return json.dumps(v, ensure_ascii=False)
    return v


def rows_for(doc, key, section, cols):
    data = doc.get(section)
    items = data if isinstance(data, list) else [data or {}]
    out = []
    for it in items:
        it = dict(it); it["key"] = key
        out.append([cell(it.get(c, it.get(c.replace(" ", "_")))) for c in cols])
    return out


def main(xlsx, keys):
    wb = load_workbook(xlsx) if os.path.exists(xlsx) else Workbook()
    if not os.path.exists(xlsx):
        wb.remove(wb.active)
    for key in keys:
        doc = json.load(open(os.path.join(RESEARCH, f"{key}.json"), encoding="utf-8"))
        status = doc["exercise"].get("status")
        if status != "owner-approved":
            sys.exit(f"{key}: status is '{status}', not 'owner-approved' — the owner has not said save it")
        for name, section, cols in SHEETS:
            ws = wb[name] if name in wb.sheetnames else wb.create_sheet(name)
            if ws.max_row == 1 and ws.cell(1, 1).value is None:
                for i, c in enumerate(cols, 1):   # ws.append would land on row 2 once A1 has been read
                    ws.cell(1, i, c)
                for c in ws[1]:
                    c.font = HEAD; c.fill = FILL
                ws.freeze_panes = "B2"
            keep = [r for r in ws.iter_rows(min_row=2, values_only=True) if r[0] != key]
            ws.delete_rows(2, ws.max_row)
            for r in keep + rows_for(doc, key, section, cols):
                ws.append(list(r))
            for i, c in enumerate(cols, 1):
                ws.column_dimensions[ws.cell(1, i).column_letter].width = 14 if c in ("key", "tier", "glow", "grade") else 32
            for row in ws.iter_rows(min_row=2):
                for c in row:
                    c.alignment = Alignment(wrap_text=True, vertical="top")
        print(f"{key}: saved")
    wb.save(xlsx)


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2:])
